"""Exporta o resultado em Excel (.xlsx), CSV e JSON.

Cabeçalhos e textos em português; chaves do JSON em inglês (snake_case), iguais às do Supabase.
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .contacts import CONTACT_LABELS
from .pipeline import FILTER_LABELS, MODE_LABELS, ORDER_LABELS, RunResult, SearchConfig, niche_summary, parse_datetime

SOURCE_NAME = "YouTube Data API v3"
SOURCE_NOTE = f"Fonte: {SOURCE_NAME}."
RETENTION_NOTE = (
    "Pelas políticas da API do YouTube, dados obtidos pela API devem ser atualizados ou apagados em até "
    "30 dias. Rode a busca de novo para atualizar e descarte arquivos antigos."
)
CONTACT_NOTE = (
    "Os contatos são só os que o próprio criador publicou na descrição do canal ou repetiu nas descrições "
    "dos vídeos. Use em abordagem comercial pontual, identifique-se e atenda pedidos de exclusão (LGPD). "
    "O e-mail comercial da aba Sobre fica atrás de um captcha: confira manualmente."
)
STATUS_OPTIONS = ("Novo", "Contatado", "Respondeu", "Em negociação", "Fechado", "Descartado")
LOCAL_TZ = ZoneInfo("America/Sao_Paulo")

_FONT = "Arial"
_HEADER_FILLS = {
    "id": "0B4F44",
    "contact": "0E7C66",
    "about": "0E7C66",
    "manual": "8A5A00",
    "metric": "2F4A45",
    "meta": "55625F",
}
_MANUAL_FILL = "FFF4CC"
_SECTION_FILL = "EEF3F1"
_LINK_COLOR = "0B5CAD"
_NUMBER_FORMATS = {
    "int": "#,##0",
    "float": "#,##0.0",
    "pct": "0.00%",
    "date": "dd/mm/yyyy",
    "datetime": "dd/mm/yyyy hh:mm",
}
_CSV_RISKY_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_MAX_CELL_CHARS = 32_000
_CONTACT_WIDTHS = {
    "whatsapp": 18, "phones": 18, "instagram": 22, "tiktok": 20, "twitter_x": 18, "facebook": 28,
    "linkedin": 30, "kwai": 16, "threads": 18, "telegram": 20, "discord": 22, "twitch": 18,
    "link_in_bio": 26, "websites": 28,
}


@dataclass(frozen=True)
class Column:
    header: str
    get: Callable[[dict[str, Any]], Any]
    width: float = 14
    kind: str = "text"  # text | int | float | pct | date | datetime
    group: str = "id"  # id | contact | about | manual | metric | meta
    note: str = ""
    link: Callable[[dict[str, Any]], str | None] | None = None
    free_text: bool = False  # texto livre de terceiros: recebe proteção contra fórmulas no CSV


def _join(values: list[str] | None) -> str:
    return ", ".join(values or [])


def _contact(field: str) -> Callable[[dict[str, Any]], str]:
    return lambda record: _join(record["contacts"].get(field))


def channel_columns(with_contacts: bool) -> list[Column]:
    """Colunas da aba Canais: identificação, contatos, aba Sobre, colunas manuais, métricas e metadados."""
    columns = [
        Column("Canal", lambda r: r["title"], 32, link=lambda r: r["url"], free_text=True),
        Column("Handle", lambda r: r["handle"], 22),
        Column("Inscritos", lambda r: r["subscribers"], 12, "int",
               note="Número arredondado pela própria API. Em branco quando o canal oculta os inscritos."),
        Column("País", lambda r: r["country"], 7),
        Column("Encontrado por", lambda r: _join(r["found_by"]), 28, free_text=True,
               note='Nichos em que o canal apareceu, ou "perfil informado".'),
    ]
    if with_contacts:
        columns += [
            Column("E-mail publicado", lambda r: r["primary_email"], 30, group="contact",
                   note="Primeiro e-mail publicado pelo criador na descrição do canal ou repetido nos vídeos."),
            Column("Outros e-mails", lambda r: _join(r["contacts"]["emails"][1:]), 26, group="contact"),
        ]
        columns += [
            Column(CONTACT_LABELS[field], _contact(field), width, group="contact")
            for field, width in _CONTACT_WIDTHS.items()
        ]
    columns += [
        Column("Aba Sobre", lambda r: "Abrir", 10, group="about", link=lambda r: r["about_url"],
               note="Abre a aba Sobre do canal. O e-mail comercial fica atrás de um captcha: "
                    "confira ali e anote na coluna ao lado."),
        Column("E-mail comercial verificado", lambda r: None, 30, group="manual",
               note="Preenchimento manual, depois de conferir na aba Sobre."),
        Column("Status", lambda r: STATUS_OPTIONS[0], 15, group="manual",
               note="Escolha na lista: " + ", ".join(STATUS_OPTIONS) + "."),
        Column("Responsável", lambda r: None, 16, group="manual"),
        Column("Observações", lambda r: None, 36, group="manual"),
        Column("Visualizações totais", lambda r: r["total_views"], 16, "int", "metric"),
        Column("Vídeos no canal", lambda r: r["video_count"], 11, "int", "metric"),
        Column("Vídeos analisados", lambda r: r["recent_videos_analyzed"], 11, "int", "metric",
               note="Vídeos recentes usados nas médias à direita."),
        Column("Média de visualizações", lambda r: r["avg_views_recent"], 14, "int", "metric",
               note="Média simples dos vídeos analisados."),
        Column("Média de curtidas", lambda r: r["avg_likes_recent"], 12, "int", "metric",
               note="Média dos vídeos analisados que exibem curtidas."),
        Column("Média de comentários", lambda r: r["avg_comments_recent"], 13, "int", "metric",
               note="Média dos vídeos analisados com comentários abertos."),
        Column("Engajamento", lambda r: r["engagement_rate_recent"], 13, "pct", "metric",
               note="(curtidas + comentários) ÷ visualizações, somando os vídeos analisados que exibem curtidas."),
        Column("Último upload", lambda r: r["last_upload_at"], 13, "date", "metric"),
        Column("Dias desde o último upload", lambda r: r["days_since_last_upload"], 12, "int", "metric"),
        Column("Intervalo médio entre uploads (dias)", lambda r: r["avg_days_between_uploads"], 14, "float",
               "metric"),
        Column("Idioma", lambda r: r["default_language"], 8, group="meta"),
        Column("Canal criado em", lambda r: r["channel_created_at"], 13, "date", "meta"),
        Column("Tópicos (YouTube)", lambda r: _join(r["topics"]), 30, group="meta",
               note="Categorias atribuídas pelo próprio YouTube."),
        Column("Palavras-chave do canal", lambda r: _join(r["keywords"]), 40, group="meta", free_text=True,
               note="Palavras-chave declaradas pelo criador nas configurações do canal."),
        Column("Ocorrências na busca", lambda r: r["search_hits"], 12, "int", "meta",
               note="Quantos resultados da busca apontaram para o canal."),
        Column("Origem", lambda r: r["source"], 22, group="meta"),
        Column("Descrição do canal", lambda r: r["description"], 60, group="meta", free_text=True),
        Column("URL do canal", lambda r: r["url"], 44, group="meta", link=lambda r: r["url"]),
        Column("ID do canal", lambda r: r["channel_id"], 27, group="meta"),
        Column("Coletado em", lambda r: r["collected_at"], 17, "datetime", "meta"),
    ]
    return columns


# ------------------------------------------------------------ tabelas (UI e CSV)
def _table_columns(result: RunResult, pct_as_percent: bool) -> list[tuple[str, Column]]:
    pairs = []
    for column in channel_columns(result.config.extract_contacts):
        if column.group == "manual":
            continue
        header = f"{column.header} (%)" if column.kind == "pct" and pct_as_percent else column.header
        pairs.append((header, column))
    return pairs


def table_rows(result: RunResult, pct_as_percent: bool = False) -> list[dict[str, Any]]:
    """Linhas planas da aba Canais, sem as colunas manuais. Datas em texto ISO, horário de Brasília."""
    pairs = _table_columns(result, pct_as_percent)
    rows = []
    for record in result.records:
        row: dict[str, Any] = {}
        for header, column in pairs:
            value = column.get(record)
            if column.group == "about":
                value = record["about_url"]
            elif column.kind == "pct" and pct_as_percent and value is not None:
                value = round(value * 100, 2)
            elif column.kind in ("date", "datetime"):
                value = _text_datetime(value, column.kind)
            row[header] = value
        rows.append(row)
    return rows


def contact_table(result: RunResult) -> list[dict[str, str]]:
    return [
        {
            "Canal": row["title"],
            "Handle": row["handle"],
            "Tipo": row["type_label"],
            "Contato": row["value"],
            "Onde foi encontrado": row["origin"],
            "URL do canal": row["url"],
        }
        for row in result.contact_rows
    ]


def niche_tables(result: RunResult) -> dict[str, list[dict[str, Any]]]:
    summary = niche_summary(result.records)
    with_contacts = result.config.extract_contacts
    niches = []
    for row in summary["niches"]:
        item: dict[str, Any] = {"Nicho": row["niche"], "Canais": row["channels"], "Inscritos (soma)": row["subscribers_sum"]}
        if with_contacts:
            item["Com e-mail"] = row["with_email"]
            item["Com contato direto"] = row["with_direct_contact"]
        niches.append(item)
    return {
        "Por nicho": niches,
        "Palavras-chave declaradas pelos canais": [
            {"Palavra-chave": row["keyword"], "Canais": row["channels"]} for row in summary["keywords"]
        ],
        "Tópicos do YouTube": [{"Tópico": row["topic"], "Canais": row["channels"]} for row in summary["topics"]],
    }


def legend(cfg: SearchConfig) -> list[tuple[str, str]]:
    items = [
        ("Inscritos", "Valor arredondado pela própria API. Em branco quando o canal oculta o número."),
        ("Médias e engajamento",
         "Contas simples sobre os vídeos recentes analisados: médias de visualizações, curtidas e comentários, "
         "e (curtidas + comentários) ÷ visualizações. Não há pontuação nem ranking de canais."),
        ("Encontrado por e ocorrências",
         "Em quais nichos o canal apareceu e quantos resultados da busca apontaram para ele. "
         "A lista vem ordenada por inscritos."),
        ("Tópicos e palavras-chave",
         "Tópicos são categorias atribuídas pelo YouTube; palavras-chave são as declaradas pelo criador. "
         "A ferramenta não classifica canais."),
        ("Colunas amarelas",
         "Para a sua equipe preencher: e-mail conferido na aba Sobre, status do contato, responsável e observações."),
    ]
    if cfg.extract_contacts:
        items.append((
            "Contatos publicados",
            "Lidos da descrição do canal e das descrições dos vídeos recentes. Links presentes em menos de "
            f"{cfg.min_video_repeats} vídeos (em geral de patrocinadores) ficam de fora.",
        ))
    return items


def summary_items(result: RunResult) -> list[tuple[str, str, str]]:
    """(seção, rótulo, valor) com parâmetros, resultado, cota, legenda e regras de uso."""
    cfg, stats = result.config, result.stats
    usage, quota = stats["usage_this_run"], stats["quota_today"]
    items: list[tuple[str, str, str]] = []

    def add(section: str, label: str, value: str) -> None:
        items.append((section, label, value))

    add("Execução", "Coletado em", f"{_local(result.finished_at):%d/%m/%Y %H:%M} (horário de Brasília)")
    add("Execução", "Fonte", f"{SOURCE_NAME}, dados públicos, sem login")
    add("Execução", "Nichos", ", ".join(cfg.niches) or "—")
    add("Execução", "Perfis informados", ", ".join(cfg.profiles) or "—")
    if cfg.niches:
        add("Execução", "Modo de busca", MODE_LABELS.get(cfg.mode, cfg.mode))
        add("Execução", "Páginas por nicho", str(cfg.pages_per_niche))
        add("Execução", "Região e idioma", f"{cfg.region or 'qualquer região'} · {cfg.language or 'qualquer idioma'}")
        add("Execução", "Ordenação da busca", ORDER_LABELS.get(cfg.order, cfg.order))
        add("Execução", "Publicados após",
            cfg.published_after.strftime("%d/%m/%Y") if cfg.published_after else "sem filtro")
    add("Execução", "Faixa de inscritos", _range_label(cfg.min_subscribers, cfg.max_subscribers))
    add("Execução", "Atividade",
        f"upload nos últimos {cfg.active_within_days} dias" if cfg.active_within_days else "sem filtro")
    add("Execução", "Limite de canais", format_int(cfg.max_channels) if cfg.max_channels else "sem limite")
    add("Execução", "Vídeos analisados por canal", str(cfg.videos_per_channel))
    if cfg.extract_contacts:
        add("Execução", "Contatos publicados", "extraídos" + (", só canais com contato" if cfg.only_with_contact else ""))
    else:
        add("Execução", "Contatos publicados", "não extraídos")

    add("Resultado", "Canais encontrados na busca", format_int(stats["channels_discovered"]))
    add("Resultado", "Canais no relatório", format_int(stats["channels_in_report"]))
    if cfg.extract_contacts:
        add("Resultado", "Com e-mail publicado", format_int(stats["with_email"]))
        add("Resultado", "Com contato direto (e-mail, WhatsApp ou telefone)", format_int(stats["with_direct_contact"]))
    if stats["unresolved_profiles"]:
        add("Resultado", "Perfis não encontrados", ", ".join(stats["unresolved_profiles"]))
    add("Resultado", "Coleta completa", "não, interrompida pela cota (veja os avisos)" if stats["interrupted"] else "sim")
    for key, count in stats["filtered"].items():
        if count:
            add("Filtrados", FILTER_LABELS.get(key, key), format_int(count))
    for warning in stats["warnings"]:
        add("Avisos", "Aviso", warning)

    add("Cota", "Buscas nesta execução", format_int(usage.get("search", 0)))
    add("Cota", "Unidades nesta execução", format_int(usage.get("units", 0)))
    add("Cota", "Respostas reaproveitadas do cache", format_int(stats["cache_hits"]))
    add("Cota", "Buscas usadas hoje", f"{format_int(quota['search']['used'])} de {format_int(quota['search']['limit'])}")
    add("Cota", "Unidades usadas hoje", f"{format_int(quota['units']['used'])} de {format_int(quota['units']['limit'])}")
    add("Cota", "Renovação", "meia-noite no horário do Pacífico (4h ou 5h em Brasília)")

    for label, text in legend(cfg):
        add("Como ler", label, text)
    add("Regras de uso", "Fonte", SOURCE_NOTE)
    add("Regras de uso", "Retenção", RETENTION_NOTE)
    if cfg.extract_contacts:
        add("Regras de uso", "Contatos", CONTACT_NOTE)
    return items


# ------------------------------------------------------------------------ Excel
def xlsx_bytes(result: RunResult) -> bytes:
    workbook = Workbook()
    workbook.properties.creator = "YT Prospector"
    workbook.properties.title = "Canais do YouTube por nicho"
    channels = workbook.active
    channels.title = "Canais"
    _channels_sheet(channels, result)
    if result.config.extract_contacts:
        _contacts_sheet(workbook.create_sheet("Contatos"), result)
    _niches_sheet(workbook.create_sheet("Nichos"), result)
    _summary_sheet(workbook.create_sheet("Resumo"), result)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _channels_sheet(ws: Worksheet, result: RunResult) -> None:
    columns = channel_columns(result.config.extract_contacts)
    _write_header(ws, columns)
    for row, record in enumerate(result.records, 2):
        for index, column in enumerate(columns, 1):
            cell = ws.cell(row=row, column=index)
            _set(cell, column.get(record), column.kind)
            target = column.link(record) if column.link else None
            if target and cell.value is not None:
                cell.hyperlink = target
                cell.font = _font(color=_LINK_COLOR, underline="single")
            if column.group == "manual":
                cell.fill = _fill(_MANUAL_FILL)
    last_row = max(len(result.records) + 1, 2)
    if not result.records:
        _set(ws.cell(row=2, column=1), "Nenhum canal passou pelos filtros. Amplie a busca ou afrouxe os filtros.")
        ws.cell(row=2, column=1).font = _font(italic=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"

    status = get_column_letter(next(i for i, column in enumerate(columns, 1) if column.header == "Status"))
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(STATUS_OPTIONS) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Status inválido",
        error="Escolha um status da lista.",
    )
    ws.add_data_validation(validation)
    validation.add(f"{status}2:{status}{last_row}")


def _contacts_sheet(ws: Worksheet, result: RunResult) -> None:
    widths = {"Canal": 32, "Handle": 22, "Tipo": 14, "Contato": 34, "Onde foi encontrado": 34, "URL do canal": 44}
    columns = [Column(header, lambda r, h=header: r[h], width) for header, width in widths.items()]
    _write_header(ws, columns)
    rows = contact_table(result)
    for row_index, row in enumerate(rows, 2):
        for index, header in enumerate(widths, 1):
            cell = ws.cell(row=row_index, column=index)
            _set(cell, row[header])
            if header in ("Canal", "URL do canal") and row["URL do canal"]:
                cell.hyperlink = row["URL do canal"]
                cell.font = _font(color=_LINK_COLOR, underline="single")
    if not rows:
        _set(ws.cell(row=2, column=1), "Nenhum contato publicado foi encontrado nos canais do relatório.")
        ws.cell(row=2, column=1).font = _font(italic=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{max(len(rows) + 1, 2)}"


def _niches_sheet(ws: Worksheet, result: RunResult) -> None:
    tables = niche_tables(result)
    start = 1
    tallest = 0
    for title, rows in tables.items():
        headers = list(rows[0]) if rows else _empty_headers(title, result.config.extract_contacts)
        title_cell = ws.cell(row=1, column=start, value=title)
        title_cell.font = _font(bold=True, size=11)
        for offset, header in enumerate(headers):
            cell = ws.cell(row=2, column=start + offset, value=header)
            cell.font = _font(bold=True, color="FFFFFF")
            cell.fill = _fill(_HEADER_FILLS["id"])
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(start + offset)].width = 34 if offset == 0 else 14
        for row_index, row in enumerate(rows, 3):
            for offset, header in enumerate(headers):
                _set(ws.cell(row=row_index, column=start + offset), row[header], "text" if offset == 0 else "int")
        if not rows:
            _set(ws.cell(row=3, column=start), "Sem dados nesta coleta.")
            ws.cell(row=3, column=start).font = _font(italic=True)
        tallest = max(tallest, len(rows) or 1)
        start += len(headers) + 1
    ws.row_dimensions[2].height = 30
    note = ws.cell(row=tallest + 4, column=1, value="Contagens simples sobre os canais deste relatório. " + SOURCE_NOTE)
    note.font = _font(italic=True, color="55625F")


def _summary_sheet(ws: Worksheet, result: RunResult) -> None:
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 100
    ws.cell(row=1, column=1, value="Resumo da coleta").font = _font(bold=True, size=13)
    row = 3
    current = None
    for section, label, value in summary_items(result):
        if section != current:
            if current is not None:
                row += 1
            for column in (1, 2):
                ws.cell(row=row, column=column).fill = _fill(_SECTION_FILL)
            ws.cell(row=row, column=1, value=section).font = _font(bold=True)
            current = section
            row += 1
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = _font()
        label_cell.alignment = Alignment(vertical="top", wrap_text=True)
        value_cell = ws.cell(row=row, column=2)
        _set(value_cell, value)
        value_cell.alignment = Alignment(vertical="top", wrap_text=True)
        row += 1


def _write_header(ws: Worksheet, columns: list[Column]) -> None:
    for index, column in enumerate(columns, 1):
        cell = ws.cell(row=1, column=index, value=column.header)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = _fill(_HEADER_FILLS[column.group])
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if column.note:
            comment = Comment(column.note, "YT Prospector")
            comment.width, comment.height = 320, 110
            cell.comment = comment
        ws.column_dimensions[get_column_letter(index)].width = column.width
    ws.row_dimensions[1].height = 42


def _empty_headers(title: str, with_contacts: bool) -> list[str]:
    if title == "Por nicho":
        return ["Nicho", "Canais", "Inscritos (soma)"] + (["Com e-mail", "Com contato direto"] if with_contacts else [])
    if title.startswith("Palavras"):
        return ["Palavra-chave", "Canais"]
    return ["Tópico", "Canais"]


def _set(cell: Any, value: Any, kind: str = "text") -> None:
    """Grava o valor com fonte, formato e proteção: texto que começa com '=' nunca vira fórmula."""
    if kind in ("date", "datetime"):
        value = _excel_datetime(value, kind)
    elif isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:_MAX_CELL_CHARS] or None
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    if value is not None and kind in _NUMBER_FORMATS:
        cell.number_format = _NUMBER_FORMATS[kind]
    cell.font = _font()


def _font(size: float = 10, **kwargs: Any) -> Font:
    return Font(name=_FONT, size=size, **kwargs)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


# ---------------------------------------------------------------- CSV e JSON
def csv_bytes(result: RunResult) -> bytes:
    """CSV (UTF-8 com BOM, separador vírgula) para importar em CRM ou planilhas online."""
    pairs = _table_columns(result, pct_as_percent=True)
    headers = [header for header, _ in pairs]
    free_text = {header for header, column in pairs if column.free_text}
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(headers)
    for row in table_rows(result, pct_as_percent=True):
        writer.writerow([_csv_value(row.get(header), header in free_text) for header in headers])
    return buffer.getvalue().encode("utf-8-sig")


def _csv_value(value: Any, free_text: bool) -> Any:
    if value is None:
        return ""
    if isinstance(value, str) and free_text and value.startswith(_CSV_RISKY_PREFIXES):
        return "'" + value  # evita que planilhas executem o texto como fórmula
    return value


def json_payload(result: RunResult) -> dict[str, Any]:
    return {
        "source": SOURCE_NAME,
        "notice": RETENTION_NOTE,
        "generated_at": result.finished_at.isoformat(),
        "config": result.config.to_dict(),
        "stats": result.stats,
        "niche_summary": niche_summary(result.records),
        "channels": result.records,
    }


def json_bytes(result: RunResult) -> bytes:
    return json.dumps(json_payload(result), ensure_ascii=False, indent=2, default=str).encode("utf-8")


# ---------------------------------------------------------------- arquivos
def default_basename(result: RunResult) -> str:
    return f"yt-prospector-{_local(result.finished_at):%Y%m%d-%H%M%S}"


def write_xlsx(result: RunResult, path: str | Path) -> Path:
    return _write(path, xlsx_bytes(result))


def write_csv(result: RunResult, path: str | Path) -> Path:
    return _write(path, csv_bytes(result))


def write_json(result: RunResult, path: str | Path) -> Path:
    return _write(path, json_bytes(result))


def _write(path: str | Path, data: bytes) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    return target


# ------------------------------------------------------------------- apoio
def format_int(value: int | float | None) -> str:
    if value is None:
        return "—"
    return f"{round(value):,}".replace(",", ".")


def _range_label(minimum: int | None, maximum: int | None) -> str:
    if minimum and maximum:
        return f"de {format_int(minimum)} a {format_int(maximum)}"
    if minimum:
        return f"a partir de {format_int(minimum)}"
    if maximum:
        return f"até {format_int(maximum)}"
    return "sem filtro"


def _local(moment: Any) -> Any:
    parsed = parse_datetime(moment)
    return parsed.astimezone(LOCAL_TZ) if parsed else None


def _excel_datetime(value: Any, kind: str) -> Any:
    local = _local(value)
    if local is None:
        return None
    naive = local.replace(tzinfo=None)  # o Excel não guarda fuso horário
    return naive.date() if kind == "date" else naive


def _text_datetime(value: Any, kind: str) -> str | None:
    local = _local(value)
    if local is None:
        return None
    return f"{local:%Y-%m-%d}" if kind == "date" else f"{local:%Y-%m-%d %H:%M}"
