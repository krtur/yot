from __future__ import annotations

import csv
import io
import json

import pytest
from openpyxl import load_workbook

import fakes
from fakes import NICHE_A, NICHE_B, NOW, cid, main_config
from prospector import export, supabase_sink
from prospector.pipeline import SearchConfig, run

MAIN_TITLES = ["Ana Finanças", "Investe Já", "Canal Direto"]


def workbook(result):
    return load_workbook(io.BytesIO(export.xlsx_bytes(result)))


def header_index(ws) -> dict[str, int]:
    return {cell.value: cell.column for cell in ws[1]}


# ------------------------------------------------------------------------ Excel
def test_xlsx_structure_formats_and_links(result):
    wb = workbook(result)
    assert wb.sheetnames == ["Canais", "Contatos", "Nichos", "Resumo"]
    ws = wb["Canais"]
    col = header_index(ws)
    for header in ("Canal", "E-mail publicado", "Aba Sobre", "E-mail comercial verificado", "Status", "Engajamento"):
        assert header in col
    assert [ws.cell(row, col["Canal"]).value for row in (2, 3, 4)] == MAIN_TITLES
    assert ws.cell(2, col["Canal"]).hyperlink.target == "https://www.youtube.com/@anafinancas"
    assert ws.cell(2, col["Aba Sobre"]).hyperlink.target == "https://www.youtube.com/@anafinancas/about"
    assert ws.cell(2, col["Inscritos"]).number_format == "#,##0"
    engagement = ws.cell(2, col["Engajamento"])
    assert engagement.number_format == "0.00%" and engagement.value == pytest.approx(0.058)
    assert ws.cell(2, col["Último upload"]).number_format == "dd/mm/yyyy"
    assert ws.cell(2, col["Status"]).value == "Novo"
    assert ws.freeze_panes == "B2" and ws.auto_filter.ref
    validations = ws.data_validations.dataValidation
    assert len(validations) == 1 and "Em negociação" in validations[0].formula1
    fonts = {cell.font.name for sheet in wb.worksheets for row in sheet.iter_rows() for cell in row if cell.value}
    assert fonts == {"Arial"}


def test_untrusted_text_never_becomes_a_formula(make_client, monkeypatch):
    snippet = fakes.CHANNELS[cid(1)]["snippet"]
    monkeypatch.setitem(snippet, "title", '=HYPERLINK("http://mal.example","clique")')
    monkeypatch.setitem(snippet, "description", "Texto com caractere de controle \x07 no meio")
    result = run(make_client(), main_config(), now=NOW)
    ws = workbook(result)["Canais"]
    col = header_index(ws)
    title = ws.cell(2, col["Canal"])
    assert title.value.startswith("=HYPERLINK") and title.data_type == "s"
    assert "\x07" not in ws.cell(2, col["Descrição do canal"]).value
    rows = list(csv.DictReader(io.StringIO(export.csv_bytes(result).decode("utf-8-sig"))))
    assert rows[0]["Canal"].startswith("'=")  # neutralizado para Excel e Sheets
    assert rows[0]["Inscritos"] == "120000"  # números não recebem prefixo


def test_exports_without_contacts_drop_contact_columns(make_client):
    result = run(make_client(), main_config(extract_contacts=False), now=NOW)
    wb = workbook(result)
    assert wb.sheetnames == ["Canais", "Nichos", "Resumo"]
    headers = [cell.value for cell in wb["Canais"][1]]
    assert "E-mail publicado" not in headers and "Aba Sobre" in headers and "Status" in headers
    assert all(record["contact_sources"] == [] for record in result.records)


def test_empty_result_still_exports(make_client):
    result = run(make_client(), SearchConfig(niches=["nicho sem resultados"]), now=NOW)
    assert result.records == []
    assert workbook(result)["Canais"]["A1"].value == "Canal"
    assert export.csv_bytes(result).decode("utf-8-sig").startswith("Canal,")
    assert json.loads(export.json_bytes(result))["channels"] == []


# ------------------------------------------------------------------ CSV e JSON
def test_csv_has_bom_and_all_rows(result):
    data = export.csv_bytes(result)
    assert data.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert [row["Canal"] for row in rows] == MAIN_TITLES
    assert rows[0]["E-mail publicado"] == "contato@anafinancas.com.br"


def test_json_payload_uses_english_keys(result):
    payload = json.loads(export.json_bytes(result))
    assert "YouTube Data API v3" in json.dumps(payload["source"])
    assert payload["config"]["niches"] == [NICHE_A, NICHE_B]
    assert [channel["title"] for channel in payload["channels"]] == MAIN_TITLES
    assert payload["channels"][0]["contacts"]["emails"] == ["contato@anafinancas.com.br"]
    assert payload["stats"]["usage_this_run"] == {"search": 3, "units": 8}


def test_write_files_creates_folders(tmp_path, result):
    for writer, suffix in ((export.write_xlsx, ".xlsx"), (export.write_csv, ".csv"), (export.write_json, ".json")):
        path = writer(result, tmp_path / "nova" / f"arquivo{suffix}")
        assert path.exists() and path.stat().st_size > 0


# -------------------------------------------------------------------- Supabase
class _Response:
    def __init__(self, status_code: int, body=None) -> None:
        self.status_code = status_code
        self.text = json.dumps(body) if body is not None else ""

    def json(self):
        if not self.text:
            raise ValueError("sem corpo")
        return json.loads(self.text)


class FakePostgrest:
    def __init__(self, *responses: tuple[int, dict]) -> None:
        self.calls: list[dict] = []
        self.responses = list(responses)

    def post(self, url, params=None, headers=None, data=None, timeout=None):
        self.calls.append({"url": url, "params": params, "headers": headers, "rows": json.loads(data)})
        return _Response(*self.responses.pop(0)) if self.responses else _Response(201)


def test_supabase_upsert_in_batches_without_team_columns(result, monkeypatch):
    monkeypatch.setattr(supabase_sink, "BATCH_SIZE", 2)
    session = FakePostgrest()
    sent = supabase_sink.upsert(result, "https://abc.supabase.co/", "sb_secret_123", session=session)
    assert sent == 3 and len(session.calls) == 2
    first = session.calls[0]
    assert first["url"] == "https://abc.supabase.co/rest/v1/yt_channels"
    assert first["params"] == {"on_conflict": "channel_id"}
    assert first["headers"]["apikey"] == "sb_secret_123" and "Authorization" not in first["headers"]
    assert "resolution=merge-duplicates" in first["headers"]["Prefer"]
    row = first["rows"][0]
    assert row["channel_id"] == cid(1) and row["emails"] == ["contato@anafinancas.com.br"] and row["tiktok"] == []
    assert row["collected_at"].startswith("2026-09-23T15:00")
    assert not {"status", "notes", "assignee", "verified_business_email"} & set(row)
    assert all(set(item) == set(row) for call in session.calls for item in call["rows"])


def test_supabase_legacy_jwt_key_also_goes_as_bearer(result):
    session = FakePostgrest()
    supabase_sink.upsert(result, "https://abc.supabase.co", "eyJhbGciOi.x.y", session=session)
    assert session.calls[0]["headers"]["Authorization"] == "Bearer eyJhbGciOi.x.y"


def test_supabase_without_contacts_does_not_overwrite_contact_columns(make_client):
    result = run(make_client(), main_config(extract_contacts=False), now=NOW)
    session = FakePostgrest()
    supabase_sink.upsert(result, "https://abc.supabase.co", "sb_secret_1", session=session)
    row = session.calls[0]["rows"][0]
    assert "emails" not in row and "primary_email" not in row and "contact_sources" not in row


@pytest.mark.parametrize(
    "status, body, expected",
    [
        (404, {"code": "PGRST205", "message": "Could not find the table"}, "supabase_schema.sql"),
        (401, {"code": "42501", "message": "new row violates row-level security policy"}, "secret key"),
        (400, {"code": "PGRST204", "message": "Could not find the 'kwai' column"}, "supabase_schema.sql"),
        (500, {"message": "boom"}, "HTTP 500"),
    ],
)
def test_supabase_errors_are_explained(result, status, body, expected):
    with pytest.raises(supabase_sink.SupabaseError, match=expected):
        supabase_sink.upsert(result, "https://abc.supabase.co", "sb_secret_1", session=FakePostgrest((status, body)))


def test_supabase_settings_are_required(result):
    with pytest.raises(supabase_sink.SupabaseError, match="SUPABASE_URL"):
        supabase_sink.upsert(result, "", "", session=FakePostgrest())
    assert supabase_sink.settings({"SUPABASE_URL": " https://x.supabase.co ", "SUPABASE_SECRET_KEY": "k"}) == (
        "https://x.supabase.co", "k", "yt_channels"
    )
